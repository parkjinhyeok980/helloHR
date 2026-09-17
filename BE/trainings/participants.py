import json
from io import BytesIO
from zipfile import BadZipFile

from django.db import transaction
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404
from django.views.decorators.http import require_GET, require_http_methods, require_POST
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from employees.models import Department, Employee

from .models import Training, TrainingParticipant
from .views import serialize_participant


HEADERS = {
    'employee_number': ('사번', 'employee_number'),
    'name': ('이름', 'name'),
    'department': ('부서', 'department'),
}


def validate_person(data):
    errors = {}
    cleaned = {}
    for field, labels in HEADERS.items():
        value = data.get(field)
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        value = str(value).strip() if value is not None else ''
        max_length = 50 if field == 'employee_number' else 100
        if not value or len(value) > max_length:
            errors[field] = f'{labels[0]}은(는) 1~{max_length}자로 입력해 주세요.'
        cleaned[field] = value
    return cleaned, errors


def enroll(training, person):
    department, _ = Department.objects.get_or_create(name=person['department'])
    employee, _ = Employee.objects.update_or_create(
        employee_number=person['employee_number'],
        defaults={'name': person['name'], 'department': department},
    )
    return TrainingParticipant.objects.get_or_create(training=training, employee=employee)


@require_http_methods(['GET', 'POST'])
def participant_list(request, training_id):
    training = get_object_or_404(Training, pk=training_id)
    if request.method == 'GET':
        people = training.participants.select_related('employee__department').prefetch_related('attendance_records').order_by('id')
        return JsonResponse({'results': [serialize_participant(item) for item in people]})

    try:
        data = json.loads(request.body)
    except (json.JSONDecodeError, UnicodeDecodeError):
        return JsonResponse({'errors': {'body': '올바른 JSON을 보내 주세요.'}}, status=400)
    if not isinstance(data, dict):
        return JsonResponse({'errors': {'body': '객체 형식의 데이터를 보내 주세요.'}}, status=400)
    person, errors = validate_person(data)
    if errors:
        return JsonResponse({'errors': errors}, status=400)
    with transaction.atomic():
        participant, created = enroll(training, person)
    return JsonResponse(serialize_participant(participant), status=201 if created else 200)


@require_http_methods(['PUT', 'DELETE'])
def participant_detail(request, training_id, participant_id):
    participant = get_object_or_404(
        TrainingParticipant.objects.select_related('employee__department'),
        pk=participant_id,
        training_id=training_id,
    )
    if request.method == 'DELETE':
        participant.delete()
        return HttpResponse(status=204)

    try:
        data = json.loads(request.body)
    except (json.JSONDecodeError, UnicodeDecodeError):
        return JsonResponse({'errors': {'body': '올바른 JSON을 보내 주세요.'}}, status=400)
    if not isinstance(data, dict):
        return JsonResponse({'errors': {'body': '객체 형식의 데이터를 보내 주세요.'}}, status=400)
    person, errors = validate_person(data)
    if errors:
        return JsonResponse({'errors': errors}, status=400)

    employee = participant.employee
    if Employee.objects.exclude(pk=employee.pk).filter(
        employee_number=person['employee_number']
    ).exists():
        return JsonResponse(
            {'errors': {'employee_number': '다른 사원이 사용 중인 사번입니다.'}}, status=409
        )

    with transaction.atomic():
        department, _ = Department.objects.get_or_create(name=person['department'])
        employee.employee_number = person['employee_number']
        employee.name = person['name']
        employee.department = department
        employee.save(update_fields=['employee_number', 'name', 'department'])
    return JsonResponse(serialize_participant(participant))


@require_POST
def participant_upload(request, training_id):
    training = get_object_or_404(Training, pk=training_id)
    upload = request.FILES.get('file')
    if not upload or not upload.name.lower().endswith('.xlsx'):
        return JsonResponse({'errors': {'file': '.xlsx 엑셀 파일을 선택해 주세요.'}}, status=400)
    if upload.size > 5 * 1024 * 1024:
        return JsonResponse({'errors': {'file': '파일 크기는 5MB 이하여야 합니다.'}}, status=400)

    try:
        workbook = load_workbook(upload, read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = [str(cell).strip() if cell is not None else '' for cell in next(rows, [])]
        indexes = {
            field: next((index for index, name in enumerate(header) if name in labels), None)
            for field, labels in HEADERS.items()
        }
        if any(index is None for index in indexes.values()):
            return JsonResponse(
                {'errors': {'file': '첫 행에 사번, 이름, 부서 열이 필요합니다.'}}, status=400
            )

        people = []
        errors = []
        seen = set()
        duplicate_rows = 0
        for row_number, row in enumerate(rows, start=2):
            if row_number > 1001:
                return JsonResponse({'errors': {'file': '한 번에 1,000명까지 업로드할 수 있습니다.'}}, status=400)
            if not any(cell is not None and str(cell).strip() for cell in row):
                continue
            data = {
                field: row[index] if index < len(row) else None
                for field, index in indexes.items()
            }
            person, row_errors = validate_person(data)
            if row_errors:
                errors.append(f'{row_number}행: {next(iter(row_errors.values()))}')
                continue
            if person['employee_number'] in seen:
                duplicate_rows += 1
                continue
            seen.add(person['employee_number'])
            people.append(person)
        if errors:
            return JsonResponse({'errors': {'file': ' / '.join(errors[:5])}}, status=400)
        if not people:
            return JsonResponse({'errors': {'file': '등록할 대상자가 없습니다.'}}, status=400)
    except (InvalidFileException, BadZipFile, OSError, ValueError, StopIteration):
        return JsonResponse({'errors': {'file': '엑셀 파일을 읽을 수 없습니다.'}}, status=400)
    finally:
        if 'workbook' in locals():
            workbook.close()

    added = 0
    skipped = duplicate_rows
    with transaction.atomic():
        for person in people:
            _, created = enroll(training, person)
            added += created
            skipped += not created

    participants = training.participants.select_related('employee__department').prefetch_related('attendance_records').order_by('id')
    return JsonResponse({
        'added': added,
        'skipped': skipped,
        'participants': [serialize_participant(item) for item in participants],
    })


@require_GET
def participant_template(request):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '교육 대상자'
    sheet.append(['사번', '이름', '부서'])
    sheet.column_dimensions['A'].width = 20
    sheet.column_dimensions['B'].width = 20
    sheet.column_dimensions['C'].width = 25
    sheet['A1'].number_format = '@'
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="training_participants_template.xlsx"'
    return response
